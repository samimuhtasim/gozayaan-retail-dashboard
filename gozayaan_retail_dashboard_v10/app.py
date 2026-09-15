
from __future__ import annotations

import asyncio
import functools
import io
import json
import os
import re
import threading
import time
from pathlib import Path
from concurrent.futures import ThreadPoolExecutor, as_completed
from contextlib import asynccontextmanager
from datetime import date, datetime, timedelta, timezone
from datetime import time as clock_time
from zoneinfo import ZoneInfo
from typing import Dict, List, Optional

import numpy as np
import pandas as pd
import requests
from fastapi import FastAPI, HTTPException
from fastapi.responses import FileResponse, JSONResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles
import uvicorn
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_CENTER
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer


# ============================================================
# Configuration
# ============================================================

BASE_DIR = Path(__file__).resolve().parent
DATA_DIR = BASE_DIR / "data"

SNAPSHOT_DASH = DATA_DIR / "Dashboard (1).xlsx"
SNAPSHOT_CSS = DATA_DIR / "CSS Helper (1).xlsx"

APPS_SCRIPT_URL = os.getenv("GOZAAYAN_SHEETS_API_URL", "https://script.google.com/macros/s/AKfycbxzvfVm3zjo8ADZBCxE7jcSnKIaISEmRzbsIfWiIAUsiHHTDnLU6uhsYlhrfYMjVYLT/exec").strip()
LIVE_REQUIRED = bool(APPS_SCRIPT_URL)
BDT = ZoneInfo("Asia/Dhaka")

# WBR weeks run Monday-Sunday. A completed week only becomes "the"
# reviewed week on the following Monday at this hour (Dhaka time) — a
# data-entry buffer so numbers can be finalised before the review locks
# in. GOZAAYAN_WBR_LOCK_HOUR lets this be changed without a code deploy.
WBR_LOCK_HOUR = int(os.getenv("GOZAAYAN_WBR_LOCK_HOUR", "17"))

REFRESH_MINUTES = int(os.getenv("GOZAAYAN_REFRESH_MINUTES", "30"))

# Apps Script throttles concurrent hits to a single deployment and can
# invalidate its redirect content key under load. Low concurrency plus a
# retry is what keeps the refresh reliable.
WARM_WORKERS = int(os.getenv("GOZAAYAN_WARM_WORKERS", "2"))
FETCH_ATTEMPTS = int(os.getenv("GOZAAYAN_FETCH_ATTEMPTS", "3"))
FETCH_BACKOFF_SECONDS = float(os.getenv("GOZAAYAN_FETCH_BACKOFF", "1.5"))
# Set GOZAAYAN_BATCH_FETCH=1 only AFTER redeploying the Apps Script
# with the ?sheets= handler. Falls back to per-tab fetching if it fails.
BATCH_FETCH = os.getenv("GOZAAYAN_BATCH_FETCH", "0").strip() == "1"
SNAPSHOT_HOUR = int(os.getenv("GOZAAYAN_SNAPSHOT_HOUR", "23"))
SNAPSHOT_MINUTE = int(os.getenv("GOZAAYAN_SNAPSHOT_MINUTE", "55"))
SNAPSHOT_DIR = BASE_DIR / "snapshots"
SNAPSHOT_DIR.mkdir(exist_ok=True)

BRANCH_SOURCE = {
    "Banani": {
        "css": "DS1Banani CSS",
        "fit": "DS1Banani FIT",
        "visa": "AM",
        "dv": "A",
    },
    "Chattogram": {
        "css": "DS1CTG CSS",
        "fit": "DS1CTG FIT",
        "visa": "AJ",
        "dv": "B",
    },
    "HQ": {
        "css": "DS1HQ CSS",
        "fit": None,
        "visa": None,
        "dv": "D",
    },
    "Motijheel": {
        "css": "DS1Mot CSS",
        "fit": "DS1Mot FIT",
        "visa": "AJ",
        "dv": "C",
    },
}

TARGET_ROWS = {
    "Banani": 19,
    "Chattogram": 45,
    "HQ": 73,
    "Motijheel": 102,
}


# ============================================================
# Utilities
# ============================================================

def num(value) -> float:
    if value is None:
        return 0.0
    if isinstance(value, (int, float, np.number)):
        if np.isfinite(float(value)):
            return float(value)
        return 0.0
    s = str(value).strip().replace(",", "")
    if not s or s.lower() in {"nan", "none", "n/a", "na", "-"}:
        return 0.0
    try:
        return float(s)
    except ValueError:
        return 0.0


def fmt_date(d: date) -> str:
    return d.isoformat()


def parse_date(value: str) -> date:
    try:
        return date.fromisoformat(value)
    except Exception as exc:
        raise ValueError(f"Invalid date: {value}") from exc


def month_start(d: date) -> date:
    return d.replace(day=1)


def month_end(d: date) -> date:
    nxt = (d.replace(day=28) + timedelta(days=4)).replace(day=1)
    return nxt - timedelta(days=1)


def q3_start() -> date:
    return date(2026, 7, 1)


def q3_end() -> date:
    return date(2026, 9, 30)


def q4_start() -> date:
    return date(2026, 10, 1)


def q4_end() -> date:
    return date(2026, 12, 31)


def h2_start() -> date:
    return date(2026, 7, 1)


def h2_end() -> date:
    return date(2026, 12, 31)


def pct(actual, target):
    target = num(target)
    return (num(actual) / target) if target else None


def change(cur, prev):
    prev = num(prev)
    return ((num(cur) - prev) / prev) if prev else None


# Display-string formatting for the PDF export, which renders static
# text rather than live Excel numbers. Mirrors the XLSX number_format
# codes in export_xlsx so both downloads read the same way.
def fmt_money(v) -> str:
    if v is None:
        return "N/A"
    return f"${num(v):,.2f}"


def fmt_int(v) -> str:
    if v is None:
        return "N/A"
    return f"{int(round(num(v))):,}"


def fmt_pct(v) -> str:
    if v is None:
        return "N/A"
    return f"{num(v) * 100:,.1f}%"


def fmt_text(v) -> str:
    return "" if v is None else str(v)


def col_idx(letter: str) -> int:
    n = 0
    for ch in letter.upper():
        n = n * 26 + (ord(ch) - 64)
    return n - 1


def normalise(df: pd.DataFrame) -> pd.DataFrame:
    df = df.copy()
    df.columns = [
        str(c).strip() if str(c).strip() else f"col_{i}"
        for i, c in enumerate(df.columns)
    ]
    return df


def get_col(df: pd.DataFrame, letter: str) -> pd.Series:
    idx = col_idx(letter)
    if idx >= df.shape[1]:
        return pd.Series([np.nan] * len(df), index=df.index)
    return df.iloc[:, idx]


def dates(df: pd.DataFrame, letter: str = "A") -> pd.Series:
    return pd.to_datetime(get_col(df, letter), errors="coerce").dt.date


def sumifs(df: pd.DataFrame, value_col: str, date_col: str, start: date, end: date) -> float:
    if df.empty:
        return 0.0
    d = dates(df, date_col)
    v = pd.to_numeric(get_col(df, value_col), errors="coerce").fillna(0)
    mask = (d >= start) & (d <= end)
    return float(v.loc[mask].sum())


def sumifs_daily(df: pd.DataFrame, value_col: str, date_col: str, start: date, end: date) -> pd.Series:
    """
    Like sumifs, but grouped by day instead of collapsed to one total.
    Returned index is a sparse set of python date objects (only dates
    that actually appear in df) — callers reindex over the full range
    themselves so missing days come through as 0 rather than absent.
    """
    if df.empty:
        return pd.Series(dtype=float)
    d = dates(df, date_col)
    v = pd.to_numeric(get_col(df, value_col), errors="coerce").fillna(0)
    mask = (d >= start) & (d <= end)
    if not mask.any():
        return pd.Series(dtype=float)
    return v.loc[mask].groupby(d.loc[mask]).sum()


def count_nonempty(df: pd.DataFrame, letter: str) -> int:
    if df.empty:
        return 0
    s = get_col(df, letter)
    return int(
        s.astype(str)
        .str.strip()
        .replace("nan", "")
        .ne("")
        .sum()
    )


# ============================================================
# Source adapters
# ============================================================

class SnapshotSource:
    def __init__(self, dashboard_path: Path, css_path: Path):
        self.dashboard_path = dashboard_path
        self.css_path = css_path
        self._dash = pd.ExcelFile(dashboard_path)
        self._css = pd.ExcelFile(css_path)
        self._dash_cache: Dict[str, pd.DataFrame] = {}
        self._css_cache: Dict[str, pd.DataFrame] = {}

    def dash(self, sheet: str) -> pd.DataFrame:
        if sheet not in self._dash_cache:
            self._dash_cache[sheet] = normalise(
                pd.read_excel(self._dash, sheet_name=sheet)
            )
        return self._dash_cache[sheet].copy()

    def css(self, sheet: str) -> pd.DataFrame:
        if sheet not in self._css_cache:
            self._css_cache[sheet] = normalise(
                pd.read_excel(self._css, sheet_name=sheet)
            )
        return self._css_cache[sheet].copy()


class AppsScriptSource:
    """Read live Google Sheet tabs through the deployed Apps Script JSON API."""

    def __init__(self, api_url: str):
        self.api_url = api_url.rstrip("/")
        self._cache: Dict[str, pd.DataFrame] = {}

    def _fetch_once(self, name: str) -> pd.DataFrame:
        response = requests.get(
            self.api_url,
            params={"sheet": name},
            timeout=60,
            allow_redirects=True,
            headers={"Accept": "application/json"},
        )
        response.raise_for_status()

        try:
            payload = response.json()
        except ValueError as exc:
            snippet = response.text[:300].replace("\n", " ")
            raise RuntimeError(
                f"Apps Script returned non-JSON for '{name}': {snippet}"
            ) from exc

        if not payload.get("ok"):
            raise RuntimeError(
                payload.get("error") or
                f"Apps Script failed for sheet '{name}'."
            )

        raw_rows = payload.get("rows") or []
        if not raw_rows:
            return pd.DataFrame()

        # Drop the workbook's first/header row, matching the XLSX adapter.
        return normalise(pd.DataFrame(raw_rows[1:]))

    def _fetch(self, name: str) -> pd.DataFrame:
        """
        Apps Script answers /exec with a 302 to a single-use
        script.googleusercontent.com URL. Under concurrent load Google
        sometimes invalidates that content key before we follow the
        redirect, producing a transient 404 on the second hop. Retrying
        gets a fresh redirect and almost always succeeds.
        """
        last_exc: Optional[Exception] = None

        for attempt in range(FETCH_ATTEMPTS):
            try:
                return self._fetch_once(name)
            except Exception as exc:
                last_exc = exc
                if attempt < FETCH_ATTEMPTS - 1:
                    time.sleep(FETCH_BACKOFF_SECONDS * (attempt + 1))

        raise RuntimeError(f"{name}: {last_exc}")

    def sheet(self, name: str) -> pd.DataFrame:
        if name not in self._cache:
            self._cache[name] = self._fetch(name)
        return self._cache[name].copy()

    def _fetch_batch(self, names) -> dict:
        """
        Single request for many tabs. Requires the Apps Script to be
        redeployed with the ?sheets= handler (see AppsScript_Code.gs).
        One redirect hop instead of N removes the transient-404 failure
        mode entirely.
        """
        response = requests.get(
            self.api_url,
            params={"sheets": ",".join(names)},
            timeout=120,
            allow_redirects=True,
            headers={"Accept": "application/json"},
        )
        response.raise_for_status()

        try:
            payload = response.json()
        except ValueError as exc:
            snippet = response.text[:300].replace("\n", " ")
            raise RuntimeError(
                f"Apps Script returned non-JSON for batch: {snippet}"
            ) from exc

        if not payload.get("ok"):
            raise RuntimeError(
                payload.get("error") or "Apps Script batch request failed."
            )

        return payload

    def warm(self, names) -> list:
        """
        Pre-fetch tabs. Returns a list of error strings rather than raising:
        a tab that fails here is still fetched lazily by sheet() on first
        use, so a partial warm must not discard the tabs that did succeed.
        """
        names = [n for n in names if n not in self._cache]
        if not names:
            return []

        errors = []

        if BATCH_FETCH:
            try:
                payload = self._fetch_batch(names)
                sheets = payload.get("sheets") or {}
                for name in names:
                    raw_rows = sheets.get(name)
                    if raw_rows is None:
                        errors.append(
                            f"{name}: missing from batch response"
                        )
                        continue
                    self._cache[name] = (
                        normalise(pd.DataFrame(raw_rows[1:]))
                        if raw_rows else pd.DataFrame()
                    )
                for name, msg in (payload.get("errors") or {}).items():
                    errors.append(f"{name}: {msg}")
                return errors
            except Exception as exc:
                # Batch endpoint not deployed yet, or failed. Fall through
                # to per-tab fetching rather than taking the app down.
                errors.append(f"batch fetch fell back to per-tab: {exc}")

        workers = min(WARM_WORKERS, len(names))

        with ThreadPoolExecutor(max_workers=workers) as pool:
            futures = {pool.submit(self._fetch, name): name for name in names}
            for future in as_completed(futures):
                name = futures[future]
                try:
                    self._cache[name] = future.result()
                except Exception as exc:
                    errors.append(str(exc))

        return errors

    @property
    def cached_tabs(self) -> int:
        return len(self._cache)


class SourceManager:
    """
    Keeps one source object in memory and refreshes it explicitly.

    Live mode is preferred. Snapshot mode is the fallback for local testing.
    """

    REQUIRED_TABS = [
        "Dashboard",
        "Target Helper",
        "DV",
        "DS1Banani FIT",
        "DS1CTG FIT",
        "DS1Mot FIT",
        "DS1Banani CSS",
        "DS1CTG CSS",
        "DS1HQ CSS",
        "DS1Mot CSS",
        "All Tour",
        "All Visa",
        "WBR Format",
    ]

    def __init__(self):
        self.lock = threading.RLock()
        self.source = None
        self.source_name = "UNINITIALIZED"
        self.last_refresh: Optional[datetime] = None
        self.last_error: Optional[str] = None
        self.degraded: bool = False
        self.refresh()

    def refresh(self):
        with self.lock:
            if LIVE_REQUIRED:
                try:
                    live = AppsScriptSource(APPS_SCRIPT_URL)
                    errors = live.warm(self.REQUIRED_TABS)

                    if live.cached_tabs == 0:
                        # Nothing came back at all — treat as a hard failure.
                        raise RuntimeError(
                            "Live Sheet refresh failed — " + " | ".join(errors)
                        )

                    # Partial warm is acceptable: sheet() lazily fetches any
                    # tab that missed, so we keep the tabs we did get.
                    self.source = live
                    self.last_refresh = datetime.now(BDT)

                    if errors:
                        self.degraded = True
                        self.source_name = "LIVE GOOGLE SHEET (PARTIAL)"
                        self.last_error = (
                            f"{len(errors)} tab(s) missed this refresh, will "
                            "load on demand — " + " | ".join(errors)
                        )
                    else:
                        self.degraded = False
                        self.source_name = "LIVE GOOGLE SHEET"
                        self.last_error = None
                    return

                except Exception as exc:
                    # In live mode, do NOT fall back to bundled snapshots.
                    # But DO keep serving the last good live data if we have
                    # it — stale numbers with a visible warning beat a dead
                    # dashboard in front of a manager.
                    self.last_error = str(exc)
                    self.degraded = True

                    if self.source is not None:
                        self.source_name = "LIVE GOOGLE SHEET (STALE)"
                        # last_refresh intentionally left at its old value so
                        # the UI can show how old the data actually is.
                    else:
                        self.source_name = "LIVE GOOGLE SHEET UNAVAILABLE"
                        self.last_refresh = None
                    return

            if not SNAPSHOT_DASH.exists() or not SNAPSHOT_CSS.exists():
                raise RuntimeError(
                    "No live Google Sheet URL configured and bundled snapshots are missing."
                )

            self.source = SnapshotSource(
                SNAPSHOT_DASH,
                SNAPSHOT_CSS,
            )
            self.source_name = "BUNDLED SNAPSHOT (LOCAL TEST MODE)"
            self.last_error = None
            self.last_refresh = datetime.now(BDT)

    def tab(self, name: str) -> pd.DataFrame:
        with self.lock:
            if self.source is None:
                # Surface the real reason as JSON. Previously this raised
                # AttributeError, which FastAPI returned as a plain-text
                # 500 that the frontend then tried to JSON.parse().
                raise HTTPException(
                    status_code=503,
                    detail=self.last_error or "Live source unavailable.",
                )
            if isinstance(self.source, SnapshotSource):
                return self.source.dash(name)
            return self.source.sheet(name)


SOURCE = SourceManager()


# ============================================================
# Business calculations
# ============================================================

def cache_per_refresh(fn):
    """
    Memoizes fn's result against SOURCE.last_refresh. Data only actually
    changes when a refresh completes, but branch_period / h2_performance /
    dashboard_h2 get recomputed from scratch on every request that touches
    them — including indirectly, e.g. wbr_metrics() calls h2_performance()
    on every single WBR load, and branch_period() gets called 4-8+ times
    per request across HQ's residual calc, h2_performance's quarters, and
    wbr_metrics' current/previous week.

    Keying on last_refresh means a refresh naturally invalidates the
    cache (new key, cache miss) with no explicit invalidation code, and a
    failed refresh that keeps serving stale data (see SourceManager)
    keeps serving the matching cached computations too, consistently.
    """
    cache: dict = {}

    @functools.wraps(fn)
    def wrapper(*args, **kwargs):
        key = (SOURCE.last_refresh, args, tuple(sorted(kwargs.items())))
        if key in cache:
            return cache[key]
        result = fn(*args, **kwargs)
        if len(cache) > 200:
            cache.clear()
        cache[key] = result
        return result

    wrapper.cache_clear = cache.clear
    return wrapper


def target_value(df: pd.DataFrame, row_1based: int, col: str) -> float:
    # pd.read_excel() consumes the first worksheet row as the header,
    # so workbook row N maps to dataframe row N-2.
    idx = row_1based - 2
    if idx < 0 or idx >= len(df):
        return 0.0
    return num(df.iloc[idx, col_idx(col)])


def ensure_source_available():
    if SOURCE.source is None:
        raise RuntimeError(
            SOURCE.last_error
            or "Live Google Sheet is unavailable."
        )


@cache_per_refresh
def dashboard_h2() -> dict:
    ensure_source_available()
    """
    The Dashboard sheet is the source of truth for the locked H2 page.
    We do not derive these values independently.
    """
    df = SOURCE.tab("Dashboard")

    # Actual workbook block: A4:J9.
    block = df.iloc[2:8, :10].fillna("")

    rows = []
    for row in block.itertuples(index=False, name=None):
        rows.append(list(row))

    return {
        "page": "dashboard",
        "currency": "USD",
        "title": "H-2 Achievement Dashboard — All Products",
        "rows": rows,
    }


TARGET_MONTH_COLS = {
    7: ("D", "E"),   # July
    8: ("I", "J"),   # August
    9: ("N", "O"),   # September
    # Target Helper does not have Oct/Nov/Dec blocks yet. The existing
    # pattern steps +5 columns per month (D->I->N), so the next blocks
    # would be Oct=(S,T), Nov=(X,Y), Dec=(AC,AD) — add them here once
    # those columns exist in the sheet. Do NOT guess a fallback month;
    # a wrong target silently misreports company performance.
}


def target_cols_for_month(month: int) -> tuple[str, str]:
    cols = TARGET_MONTH_COLS.get(month)
    if cols is None:
        raise HTTPException(
            status_code=409,
            detail=(
                f"Target Helper has no target columns configured for "
                f"month {month}. Add the next block to the sheet "
                "following the existing +5-column pattern, then add "
                "the entry to TARGET_MONTH_COLS in app.py."
            ),
        )
    return cols


@cache_per_refresh
def branch_month_targets(branch: str, end: date) -> dict:
    ensure_source_available()
    t = SOURCE.tab("Target Helper")
    row = TARGET_ROWS[branch]
    cols = target_cols_for_month(end.month)

    return {
        "Flight": target_value(t, row, cols[0]) + target_value(t, row, cols[1]),
        "Hotel": target_value(t, row + 1, cols[0]) + target_value(t, row + 1, cols[1]),
        "Tour": (
            target_value(t, row + 2, cols[0])
            + target_value(t, row + 2, cols[1])
            + target_value(t, row + 3, cols[0])
            + target_value(t, row + 3, cols[1])
        ),
        "Visa": target_value(t, row + 4, cols[0]) + target_value(t, row + 4, cols[1]),
    }


@cache_per_refresh
def branch_period(branch: str, start: date, end: date) -> dict:
    ensure_source_available()
    cfg = BRANCH_SOURCE[branch]

    css = SOURCE.tab(cfg["css"])
    fit = SOURCE.tab(cfg["fit"]) if cfg["fit"] else None

    # CSS Helper-fed DS1 CSS
    css_flight = sumifs(css, "F", "A", start, end)
    css_hotel = sumifs(css, "G", "A", start, end)
    css_tour = sumifs(css, "H", "A", start, end)

    css_flight_recv = sumifs(css, "L", "A", start, end)
    css_hotel_recv = sumifs(css, "M", "A", start, end)
    css_tour_recv = sumifs(css, "N", "A", start, end)

    css_flight_bookings = sumifs(css, "O", "A", start, end)
    css_hotel_bookings = sumifs(css, "P", "A", start, end)

    # Three FIT options are imported to the Dashboard workbook.
    retail_flight = retail_hotel = retail_tour = retail_visa = footfall = 0.0
    fit_flight_bookings = fit_hotel_bookings = 0.0

    if fit is not None:
        retail_flight = sumifs(fit, "J", "A", start, end)
        retail_hotel = sumifs(fit, "M", "A", start, end)
        retail_tour = sumifs(fit, "N", "A", start, end)
        retail_visa = sumifs(fit, cfg["visa"], "A", start, end)
        footfall = sumifs(fit, "Q", "A", start, end)

        fit_flight_bookings = sumifs(fit, "R", "A", start, end)
        fit_hotel_bookings = sumifs(fit, "V", "A", start, end)

    # HQ has no dedicated FIT sheet in the current workbook.
    # Its retail Tour and Visa values are residuals from the global
    # All Tour / All Visa figures after the other branches' contributions.
    if branch == "HQ":
        all_tour = SOURCE.tab("All Tour")
        all_visa = SOURCE.tab("All Visa")

        other_tour = sum(
            branch_period(b, start, end)["products"]["Tour"]["gmv"]
            for b in ["Banani", "Chattogram", "Motijheel"]
        )
        other_visa = sum(
            branch_period(b, start, end)["products"]["Visa"]["gmv"]
            for b in ["Banani", "Chattogram", "Motijheel"]
        )

        retail_tour = (
            sumifs(all_tour, "O", "A", start, end)
            - other_tour
            - css_tour
        )

        retail_visa = (
            sumifs(all_visa, "D", "A", start, end)
            - other_visa
        )

    products = {
        "Flight": {
            "css": css_flight,
            "retail": retail_flight,
            "gmv": css_flight + retail_flight,
            "receivable": css_flight_recv,
            "bookings": css_flight_bookings + fit_flight_bookings,
        },
        "Tour": {
            "css": css_tour,
            "retail": retail_tour,
            "gmv": css_tour + retail_tour,
            "receivable": css_tour_recv,
            "bookings": 0.0,
        },
        "Visa": {
            "css": 0.0,
            "retail": retail_visa,
            "gmv": retail_visa,
            "receivable": 0.0,
            "bookings": 0.0,
        },
        "Hotel": {
            "css": css_hotel,
            "retail": retail_hotel,
            "gmv": css_hotel + retail_hotel,
            "receivable": css_hotel_recv,
            "bookings": css_hotel_bookings + fit_hotel_bookings,
        },
        "Vouchers": {
            "css": 0.0,
            "retail": 0.0,
            "gmv": 0.0,
            "receivable": 0.0,
            "bookings": 0.0,
        },
        "Others": {
            "css": 0.0,
            "retail": 0.0,
            "gmv": 0.0,
            "receivable": 0.0,
            "bookings": 0.0,
        },
    }

    total = sum(p["gmv"] for p in products.values())

    # branch_period's GMV computation (above) has nothing to do with
    # whether the target month's columns exist in Target Helper yet —
    # but h2_performance() calls branch_period() for Q4 (Oct-Dec) purely
    # to sum actual GMV, every single time the H2 dashboard OR the WBR
    # page loads, months before Q4 starts and before those columns
    # exist. Letting that hard-fail here would take down pages that
    # never asked for a target. So: GMV always computes and is always
    # trustworthy; a missing target degrades to 0 with an explicit flag
    # rather than crashing the request. Callers that actually display a
    # target (branch pages, exports) can check targets_available and
    # show "not yet configured" instead of a misleading $0.
    try:
        targets = branch_month_targets(branch, end)
        targets_available = True
    except HTTPException:
        targets = {"Flight": 0.0, "Hotel": 0.0, "Tour": 0.0, "Visa": 0.0}
        targets_available = False

    days_in_range = (end - start).days + 1

    return {
        "branch": branch,
        "start": start.isoformat(),
        "end": end.isoformat(),
        "currency": "USD",
        "products": products,
        "total_gmv": total,
        "targets": targets,
        "targets_available": targets_available,
        "footfall": footfall,
        "footfall_avg_daily": (
            footfall / days_in_range if days_in_range > 0 else 0.0
        ),
    }


@cache_per_refresh
def daily_company_gmv(start: date, end: date) -> list[dict]:
    """
    Company-wide GMV per calendar day across [start, end], for the daily
    GMV dashboard tab. Every day in range is present, zero-filled if no
    records fell on it (so a chart doesn't show gaps).

    Deliberately does NOT reuse branch_period's per-branch composition
    (including HQ's tour/visa residual trick, which exists to attribute
    a company-wide total DOWN to a branch). For a company TOTAL, that
    algebra collapses exactly to the raw column sums used below:
      - Tour: HQ's retail_tour residual = All_Tour_total - other
        branches' tour - HQ's css_tour, so summing all 4 branches'
        tour always equals All_Tour_total directly. Same for Visa
        (which has no CSS component at all).
      - Flight/Hotel: no residual trick is used for these at all — every
        branch (HQ included) sources them directly from its own CSS/FIT
        columns, so summing CSS+FIT across branches IS the company
        total, per day exactly as it is per period.
    test_patches.py checks this reconciles against summing
    branch_period(...)["total_gmv"] across branches for the same range.
    """
    ensure_source_available()
    idx = list(pd.date_range(start, end, freq="D").date)

    flight = pd.Series(0.0, index=idx)
    hotel = pd.Series(0.0, index=idx)

    for cfg in BRANCH_SOURCE.values():
        css = SOURCE.tab(cfg["css"])
        flight = flight.add(sumifs_daily(css, "F", "A", start, end), fill_value=0.0)
        hotel = hotel.add(sumifs_daily(css, "G", "A", start, end), fill_value=0.0)

        if cfg["fit"]:
            fit = SOURCE.tab(cfg["fit"])
            flight = flight.add(sumifs_daily(fit, "J", "A", start, end), fill_value=0.0)
            hotel = hotel.add(sumifs_daily(fit, "M", "A", start, end), fill_value=0.0)

    tour = pd.Series(0.0, index=idx).add(
        sumifs_daily(SOURCE.tab("All Tour"), "O", "A", start, end), fill_value=0.0
    )
    visa = pd.Series(0.0, index=idx).add(
        sumifs_daily(SOURCE.tab("All Visa"), "D", "A", start, end), fill_value=0.0
    )

    rows = []
    for d in idx:
        f = float(flight.get(d, 0.0))
        h = float(hotel.get(d, 0.0))
        t = float(tour.get(d, 0.0))
        v = float(visa.get(d, 0.0))
        rows.append({
            "date": d.isoformat(),
            "flight": f,
            "hotel": h,
            "tour": t,
            "visa": v,
            "total": f + h + t + v,
        })
    return rows


def branch_progress(branch: str, as_of: date) -> dict:
    mtd = branch_period(branch, month_start(as_of), as_of)
    targets = mtd["targets"]

    days_elapsed = as_of.day
    days_in_month = month_end(as_of).day

    projected = (
        mtd["total_gmv"] / days_elapsed * days_in_month
        if days_elapsed else 0
    )

    current_run_rate = (
        mtd["total_gmv"] / days_elapsed
        if days_elapsed else 0
    )

    monthly_target = sum(targets.values())
    remaining_days = days_in_month - days_elapsed

    required = (
        max(
            0.0,
            (
                monthly_target -
                mtd["total_gmv"]
            ) / remaining_days
        )
        if remaining_days > 0
        else 0
    )

    cfg = BRANCH_SOURCE[branch]
    dv = SOURCE.tab("DV")
    active_css = count_nonempty(dv, cfg["dv"])

    result = {
        **mtd,
        "projected": projected,
        "current_run_rate": current_run_rate,
        "required_run_rate": required,
        "gap": current_run_rate - required,
        "active_css": active_css,
        "pipeline_number": 5 if branch != "HQ" else 0,
        "pipeline_worth": 25000 if branch != "HQ" else 0,
    }

    return result


@cache_per_refresh
def h2_performance() -> list[dict]:
    ensure_source_available()
    """
    H2 is intentionally derived from the same data families:
      Flight/Hotel = branch CSS + branch FIT
      Tour/Visa = All Tour / All Visa
    Targets come from the locked Dashboard values.
    """
    dash = SOURCE.tab("Dashboard")

    q3_target = {
        "Flight": target_value(dash, 5, "B"),
        "Hotel": target_value(dash, 6, "B"),
        "Tour": target_value(dash, 7, "B"),
        "Visa": target_value(dash, 8, "B"),
    }

    h2_target = {
        "Flight": target_value(dash, 5, "H"),
        "Hotel": target_value(dash, 6, "H"),
        "Tour": target_value(dash, 7, "H"),
        "Visa": target_value(dash, 8, "H"),
    }

    q3 = {}
    q4 = {}

    for product in ["Flight", "Hotel"]:
        q3[product] = sum(
            branch_period(
                b,
                q3_start(),
                q3_end()
            )["products"][product]["gmv"]
            for b in BRANCH_SOURCE
        )

        q4[product] = sum(
            branch_period(
                b,
                q4_start(),
                q4_end()
            )["products"][product]["gmv"]
            for b in BRANCH_SOURCE
        )

    tour = SOURCE.tab("All Tour")
    visa = SOURCE.tab("All Visa")

    q3["Tour"] = sumifs(tour, "O", "A", q3_start(), q3_end())
    q4["Tour"] = sumifs(tour, "O", "A", q4_start(), q4_end())

    q3["Visa"] = sumifs(visa, "D", "A", q3_start(), q3_end())
    q4["Visa"] = sumifs(visa, "D", "A", q4_start(), q4_end())

    rows = []

    for product in ["Flight", "Hotel", "Tour", "Visa"]:
        q4_target = h2_target[product] - q3_target[product]

        q3_actual = q3[product]
        q4_actual = q4[product]
        h2_actual = q3_actual + q4_actual

        rows.append({
            "product": product,
            "q3_target": q3_target[product],
            "q3_actual": q3_actual,
            "q3_achievement": pct(q3_actual, q3_target[product]),
            "q4_target": q4_target,
            "q4_actual": q4_actual,
            "q4_achievement": pct(q4_actual, q4_target),
            "h2_target": h2_target[product],
            "h2_actual": h2_actual,
            "h2_achievement": pct(h2_actual, h2_target[product]),
        })

    total = {
        "product": "Overall — All Products",
        "q3_target": sum(r["q3_target"] for r in rows),
        "q3_actual": sum(r["q3_actual"] for r in rows),
        "q4_target": sum(r["q4_target"] for r in rows),
        "q4_actual": sum(r["q4_actual"] for r in rows),
        "h2_target": sum(r["h2_target"] for r in rows),
        "h2_actual": sum(r["h2_actual"] for r in rows),
    }

    total["q3_achievement"] = pct(
        total["q3_actual"],
        total["q3_target"]
    )

    total["q4_achievement"] = pct(
        total["q4_actual"],
        total["q4_target"]
    )

    total["h2_achievement"] = pct(
        total["h2_actual"],
        total["h2_target"]
    )

    return rows + [total]


def _week_monday(d: date) -> date:
    return d - timedelta(days=d.weekday())


def wbr_rollover_cutover(monday_of_week: date) -> datetime:
    """
    The exact instant (Dhaka time) at which the week starting
    monday_of_week becomes eligible to be "locked in" as the reviewed
    week — i.e. the Monday that starts it, at WBR_LOCK_HOUR.
    """
    return datetime.combine(
        monday_of_week,
        clock_time(WBR_LOCK_HOUR, 0),
        tzinfo=BDT,
    )


def locked_wbr_sunday(now: Optional[datetime] = None) -> date:
    """
    The 'week ending' WBR should show by default, right now.

    Was: whatever date the frontend happened to send (often 'today'),
    snapped to the Sunday of ITS OWN week — so mid-week, WBR showed the
    in-progress current week's partial numbers under a 'week ending'
    label. That's the reported bug: a week that just began, displayed as
    though it were a completed review period.

    Fix: weeks run Monday-Sunday and a completed week only becomes the
    reviewed week on the following Monday at WBR_LOCK_HOUR Dhaka time.
    Before that instant, WBR keeps showing the previous locked week so
    there's a buffer to finish entering the week's data.

    Example: week ending Sun 13 Sep completes at midnight Sun->Mon. It
    becomes the reviewed week at Mon 14 Sep, WBR_LOCK_HOUR:00 Dhaka. Any
    time before that (including Monday morning), locked_wbr_sunday()
    still returns the Sunday before that (6 Sep). Any time from Monday
    17:00 through the following Monday 16:59, it returns 13 Sep.
    """
    if now is None:
        now = datetime.now(BDT)
    elif now.tzinfo is None:
        now = now.replace(tzinfo=BDT)

    this_monday = _week_monday(now.date())
    cutover = wbr_rollover_cutover(this_monday)

    candidate_monday = this_monday if now >= cutover else this_monday - timedelta(days=7)
    return candidate_monday - timedelta(days=1)


def next_wbr_rollover(now: Optional[datetime] = None) -> datetime:
    """The next instant the locked WBR week will advance."""
    if now is None:
        now = datetime.now(BDT)
    elif now.tzinfo is None:
        now = now.replace(tzinfo=BDT)

    this_monday = _week_monday(now.date())
    cutover = wbr_rollover_cutover(this_monday)
    return cutover if now < cutover else wbr_rollover_cutover(this_monday + timedelta(days=7))


def wbr_review_month(monday: date) -> int:
    """
    Which Target Helper month-block a Mon-Sun week's targets should come
    from. A week can straddle two calendar months (e.g. 31 Aug - 6 Sep);
    ISO 8601 resolves that by the month containing the week's Thursday,
    which is the convention used here too.
    """
    return (monday + timedelta(days=3)).month


@cache_per_refresh
def wbr_targets(review_month: int) -> dict:
    """
    Was hardcoded to August's Target Helper columns (I, J) unconditionally
    — so every WBR review since August rolled into September has been
    comparing this week's actuals against last month's target. Now takes
    the month actually being reviewed and looks up its columns via
    target_cols_for_month, the same lookup branch_month_targets uses, so
    the two can't drift apart again.

    Row layout per branch block (TARGET_ROWS gives each block's base row):
      base+0 Flight, base+1 Hotel, base+2 Domestic Tour,
      base+3 International Tour, base+4 Visa.
    Two columns are summed per cell (col0 + col1) — that pairing is a
    property of how Target Helper lays out each month's block, unchanged
    from the original code.
    """
    ensure_source_available()
    t = SOURCE.tab("Target Helper")
    col0, col1 = target_cols_for_month(review_month)

    def v(row: int) -> float:
        return target_value(t, row, col0) + target_value(t, row, col1)

    bases = list(TARGET_ROWS.values())  # Banani, Chattogram, HQ, Motijheel

    return {
        "Flight": sum(v(b) for b in bases) / 4,
        "Hotel": sum(v(b + 1) for b in bases) / 4,
        "Tour": sum(v(b + 2) + v(b + 3) for b in bases) / 4,
        "Visa": sum(v(b + 4) for b in bases) / 4,
    }


def wbr_metrics(week_end: date) -> dict:
    ensure_source_available()
    # Any date passed in resolves to the Sunday of the week containing
    # it — this snapping was always correct; see locked_wbr_sunday() for
    # what actually decides the DEFAULT date the caller passes.
    monday = _week_monday(week_end)
    sunday = monday + timedelta(days=6)
    review_month = wbr_review_month(monday)

    prev_monday = monday - timedelta(days=7)
    prev_sunday = sunday - timedelta(days=7)

    def period(start, end):
        flight = sum(
            branch_period(
                b,
                start,
                end
            )["products"]["Flight"]["gmv"]
            for b in BRANCH_SOURCE
        )

        hotel = sum(
            branch_period(
                b,
                start,
                end
            )["products"]["Hotel"]["gmv"]
            for b in BRANCH_SOURCE
        )

        tour = sumifs(
            SOURCE.tab("All Tour"),
            "O",
            "A",
            start,
            end
        )

        visa = sumifs(
            SOURCE.tab("All Visa"),
            "D",
            "A",
            start,
            end
        )

        # Bookings:
        flight_bookings = sum(
            sum(
                [
                    sumifs(SOURCE.tab(cfg["css"]), "O", "A", start, end),
                    sumifs(
                        SOURCE.tab(cfg["fit"]), "R", "A", start, end
                    ) if cfg["fit"] else 0,
                ]
            )
            for cfg in BRANCH_SOURCE.values()
        )

        hotel_bookings = sum(
            sum(
                [
                    sumifs(SOURCE.tab(cfg["css"]), "P", "A", start, end),
                    sumifs(
                        SOURCE.tab(cfg["fit"]), "V", "A", start, end
                    ) if cfg["fit"] else 0,
                ]
            )
            for cfg in BRANCH_SOURCE.values()
        )

        tour_bookings = sumifs(
            SOURCE.tab("All Tour"),
            "L",
            "A",
            start,
            end
        )

        visa_bookings = sumifs(
            SOURCE.tab("All Visa"),
            "F",
            "A",
            start,
            end
        )

        # Receivable
        flight_recv = sum(
            sumifs(
                SOURCE.tab(cfg["css"]),
                "L",
                "A",
                start,
                end
            )
            for cfg in BRANCH_SOURCE.values()
        )

        tour_recv = sum(
            sumifs(
                SOURCE.tab(cfg["css"]),
                "N",
                "A",
                start,
                end
            )
            for cfg in BRANCH_SOURCE.values()
        )

        hotel_recv = sum(
            sumifs(
                SOURCE.tab(cfg["css"]),
                "M",
                "A",
                start,
                end
            )
            for cfg in BRANCH_SOURCE.values()
        )

        # Net revenue follows the current WBR workbook sources.
        tour_net = sumifs(
            SOURCE.tab("All Tour"),
            "I",
            "A",
            start,
            end
        )

        visa_net = sumifs(
            SOURCE.tab("All Visa"),
            "L",
            "A",
            start,
            end
        )

        # WBR source has a fixed Hotel net revenue value.
        hotel_net = 83.5

        return {
            "flight": flight,
            "tour": tour,
            "visa": visa,
            "hotel": hotel,
            "gmv": flight + tour + visa + hotel,

            "flight_bookings": flight_bookings,
            "tour_bookings": tour_bookings,
            "visa_bookings": visa_bookings,
            "hotel_bookings": hotel_bookings,
            "bookings": (
                flight_bookings +
                tour_bookings +
                visa_bookings +
                hotel_bookings
            ),

            "flight_recv": flight_recv,
            "tour_recv": tour_recv,
            "hotel_recv": hotel_recv,
            "receivable": (
                flight_recv +
                tour_recv +
                hotel_recv
            ),

            "tour_net": tour_net,
            "visa_net": visa_net,
            "hotel_net": hotel_net,
            "net_revenue": (
                tour_net +
                visa_net +
                hotel_net
            ),
        }

    cur = period(monday, sunday)
    prev = period(prev_monday, prev_sunday)

    return {
        "start": monday.isoformat(),
        "end": sunday.isoformat(),

        "summary": {
            "gmv": {
                "Flight": cur["flight"],
                "Tour": cur["tour"],
                "Visa": cur["visa"],
                "Hotel": cur["hotel"],
                "Total": cur["gmv"],
                "WoW": change(cur["gmv"], prev["gmv"]),
            },

            "bookings": {
                "Flight": cur["flight_bookings"],
                "Tour": cur["tour_bookings"],
                "Visa": cur["visa_bookings"],
                "Hotel": cur["hotel_bookings"],
                "Total": cur["bookings"],
                "WoW": change(cur["bookings"], prev["bookings"]),
            },

            "receivable": {
                "Flight": cur["flight_recv"],
                "Tour": cur["tour_recv"],
                "Visa": 0.0,
                "Hotel": cur["hotel_recv"],
                "Total": cur["receivable"],
                "Ratio": pct(cur["receivable"], cur["gmv"]),
            },

            "net_revenue": {
                "Flight": 0.0,
                "Tour": cur["tour_net"],
                "Visa": cur["visa_net"],
                "Hotel": cur["hotel_net"],
                "Total": cur["net_revenue"],
                "Ratio": pct(cur["net_revenue"], cur["gmv"]),
            },
        },

        "weekly_targets": wbr_targets(review_month),
        "target_month": review_month,

        "h2": h2_performance(),

        "footfall": sum(
            branch_period(b, monday, sunday)["footfall"]
            for b in BRANCH_SOURCE
        ),
        "footfall_avg_daily": sum(
            branch_period(b, monday, sunday)["footfall"]
            for b in BRANCH_SOURCE
        ) / 7,

        "previous": {
            "gmv": prev["gmv"],
            "bookings": prev["bookings"],
            "receivable": prev["receivable"],
            "net_revenue": prev["net_revenue"],
        },

        "narrative": wbr_narrative(),
    }


def wbr_narrative() -> dict:
    ensure_source_available()
    df = SOURCE.tab("WBR Format")

    def rows(start, end, cols):
        result = []
        for r in range(start, end + 1):
            values = []
            for c in cols:
                idx = col_idx(c)
                row_idx = r - 2
                if (
                    0 <= row_idx < len(df)
                    and idx < df.shape[1]
                ):
                    v = df.iloc[row_idx, idx]
                    if pd.notna(v) and str(v).strip():
                        values.append(str(v))
            if values:
                result.append(values)
        return result

    return {
        "analysis": rows(31, 36, ["A", "C"]),
        "challenges": rows(39, 43, ["A", "B", "G", "I"]),
        "actions": rows(46, 51, ["A", "B", "G", "H", "I", "K"]),
    }


# ============================================================
# EXPORTS + SNAPSHOTS
# ============================================================

def now_bdt() -> datetime:
    return datetime.now(ZoneInfo("Asia/Dhaka"))


def safe_filename(s: str) -> str:
    return re.sub(r"[^A-Za-z0-9._-]+", "_", s).strip("_")


def report_title(page: str, start: Optional[str] = None, end: Optional[str] = None) -> str:
    if page == "dashboard":
        return "H2 Achievement Dashboard"
    if page == "wbr":
        return f"Weekly Business Review {start or ''} to {end or ''}".strip()
    return f"{page.title()} Sales Summary {start or ''} to {end or ''}".strip()


def build_export_payload(page: str, start: Optional[str], end: Optional[str]):
    if page == "dashboard":
        return dashboard_h2()
    if page == "wbr":
        raw = end or start
        target = parse_date(raw) if raw else locked_wbr_sunday()
        return wbr_metrics(target)
    if page.lower() in {"banani","chattogram","hq","motijheel"}:
        canonical = {
            "banani":"Banani",
            "chattogram":"Chattogram",
            "hq":"HQ",
            "motijheel":"Motijheel"
        }[page.lower()]
        sd = parse_date(start or end)
        ed = parse_date(end or start)
        return {
            "page":"branch",
            "branch":canonical,
            "currency":"USD",
            "start":sd.isoformat(),
            "end":ed.isoformat(),
            "data":branch_period(canonical, sd, ed),
            "progress":branch_progress(canonical, ed)
        }
    raise ValueError("Unknown page")


def export_xlsx(page: str, start: Optional[str], end: Optional[str]) -> bytes:
    payload = build_export_payload(page, start, end)
    wb = Workbook()
    ws = wb.active
    ws.title = "Dashboard"
    blue = "1D4DA2"
    deep = "00026E"
    yellow = "FFDF3E"

    # Excel number_format codes. Applied to the numeric cell value itself
    # (not a pre-formatted string), so the exported figure stays a real,
    # sortable/summable number in Excel — it just displays cleanly instead
    # of showing raw float precision like 18597.90312434.
    MONEY_FMT = '"$"#,##0.00'
    INT_FMT = "#,##0"
    PCT_FMT = "0.0%"  # pct()/change() return fractions (0.783), not 78.3

    ws["A1"] = report_title(page, start, end)
    ws["A1"].font = Font(size=18, bold=True, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor=deep)
    ws.merge_cells("A1:H1")

    r = 3

    def write_row(values, header=False, formats=None):
        nonlocal r
        for c, value in enumerate(values, 1):
            fmt = formats[c - 1] if formats and c - 1 < len(formats) else None
            if value is None:
                # e.g. WoW% with no prior-period base, or Q4 actuals
                # before Q4 starts. Blank-with-format read as $0.00,
                # which misreports "no data" as "zero" — write text.
                cell = ws.cell(r, c, "N/A" if fmt else None)
            else:
                cell = ws.cell(r, c, value)
                if fmt:
                    cell.number_format = fmt
            cell.alignment = Alignment(vertical="center")
            if header:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor=blue)
        r += 1

    if page == "dashboard":
        write_row(["Product","Q3 Target USD","Q3 Actual USD","Q3 Achievement",
                   "Q4 Target USD","Q4 Actual USD","Q4 Achievement",
                   "H2 Target USD","H2 Actual USD","H2 Achievement"], True)
        dash_fmt = [None, MONEY_FMT, MONEY_FMT, PCT_FMT,
                    MONEY_FMT, MONEY_FMT, PCT_FMT,
                    MONEY_FMT, MONEY_FMT, PCT_FMT]
        # payload["rows"][0] is the Dashboard sheet's OWN header row — the
        # export was including it as a data row underneath the header we
        # just wrote above, so every export showed the column labels
        # twice. dashboard_h2()'s API shape is unchanged (in case the
        # live frontend depends on rows[0] being the header); only the
        # export skips it.
        for row in payload["rows"][1:]:
            if not row:
                continue
            write_row(row, formats=dash_fmt)

    elif page == "wbr":
        s = payload["summary"]
        write_row(["Metric","Flight","Tour","Visa","Hotel","Total","Change / Ratio"], True)
        money_row = [None, MONEY_FMT, MONEY_FMT, MONEY_FMT, MONEY_FMT, MONEY_FMT, PCT_FMT]
        int_row = [None, INT_FMT, INT_FMT, INT_FMT, INT_FMT, INT_FMT, PCT_FMT]
        for key, label, pct_key, fmts in [
            ("gmv", "GMV", "WoW", money_row),
            ("bookings", "Bookings", "WoW", int_row),
            ("receivable", "Receivable", "Ratio", money_row),
            ("net_revenue", "Net Revenue", "Ratio", money_row),
        ]:
            d = s[key]
            write_row(
                [label, d.get("Flight"), d.get("Tour"), d.get("Visa"),
                 d.get("Hotel"), d.get("Total"), d.get(pct_key)],
                formats=fmts,
            )
        r += 1
        write_row(["Weekly Target vs Achievement"], True)
        write_row(["Product","Weekly Target USD","Actual USD","Achievement"], True)
        for prod, target in payload["weekly_targets"].items():
            actual = s["gmv"].get(prod, 0)
            write_row(
                [prod, target, actual, pct(actual, target)],
                formats=[None, MONEY_FMT, MONEY_FMT, PCT_FMT],
            )

    else:
        d = payload["data"]
        p = d["products"]
        write_row(["Product","CSS GMV USD","Retail GMV USD","Total GMV USD","Receivable USD","Bookings"], True)
        branch_fmt = [None, MONEY_FMT, MONEY_FMT, MONEY_FMT, MONEY_FMT, INT_FMT]
        for prod in ["Flight", "Tour", "Visa", "Hotel", "Vouchers", "Others"]:
            x = p[prod]
            write_row(
                [prod, x["css"], x["retail"], x["gmv"], x["receivable"], x["bookings"]],
                formats=branch_fmt,
            )
        r += 1
        write_row(["KPI", "Value"], True)
        write_row(["Selected Period GMV", d["total_gmv"]], formats=[None, MONEY_FMT])
        write_row(["MTD GMV", payload["progress"]["total_gmv"]], formats=[None, MONEY_FMT])
        write_row(["Projected GMV", payload["progress"]["projected"]], formats=[None, MONEY_FMT])
        write_row(["Current Run Rate", payload["progress"]["current_run_rate"]], formats=[None, MONEY_FMT])
        write_row(["Required Run Rate", payload["progress"]["required_run_rate"]], formats=[None, MONEY_FMT])
        write_row(["Active CSS", payload["progress"]["active_css"]], formats=[None, INT_FMT])
        write_row(["Pipeline Number", payload["progress"]["pipeline_number"]], formats=[None, INT_FMT])
        write_row(["Pipeline Worth USD", payload["progress"]["pipeline_worth"]], formats=[None, MONEY_FMT])

    for col, width in zip("ABCDEFGHIJ", [24,16,16,16,16,16,18,18,18,18]):
        ws.column_dimensions[col].width = width

    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()


def export_pdf(page: str, start: Optional[str], end: Optional[str]) -> bytes:
    payload = build_export_payload(page, start, end)
    bio = io.BytesIO()
    doc = SimpleDocTemplate(
        bio,
        pagesize=landscape(A4),
        rightMargin=24,leftMargin=24,topMargin=24,bottomMargin=24
    )
    styles = getSampleStyleSheet()
    title = ParagraphStyle("TitleX", parent=styles["Title"], fontSize=20, textColor=colors.HexColor("#00026E"), alignment=TA_CENTER)
    story = [Paragraph(report_title(page, start, end), title), Spacer(1, 10)]
    tiny = ParagraphStyle("TinyX", parent=styles["BodyText"], fontSize=8)

    def fmt_row(values, kinds):
        # ReportLab's Table stringifies cells with str() — a raw None
        # became the literal text "None" and a raw float showed full
        # precision (18597.90312...). Format explicitly per column kind
        # instead: text | money | int | pct.
        out = []
        for value, kind in zip(values, kinds):
            if kind == "money":
                out.append(fmt_money(value))
            elif kind == "int":
                out.append(fmt_int(value))
            elif kind == "pct":
                out.append(fmt_pct(value))
            else:
                out.append(fmt_text(value))
        return out

    if page == "dashboard":
        data = [["Product","Q3 Target","Q3 Actual","Q3 Ach.","Q4 Target","Q4 Actual","Q4 Ach.","H2 Target","H2 Actual","H2 Ach."]]
        dash_kinds = ["text","money","money","pct","money","money","pct","money","money","pct"]
        # payload["rows"][0] is the Dashboard sheet's own header row — see
        # the matching note in export_xlsx. Skipped here for the same
        # reason: this table already has its own header above.
        for row in payload["rows"][1:]:
            if row:
                data.append(fmt_row(row[:10], dash_kinds))
        table = Table(data, repeatRows=1)

    elif page == "wbr":
        s = payload["summary"]
        data = [["Metric","Flight","Tour","Visa","Hotel","Total","Change / Ratio"]]
        money_kinds = ["text","money","money","money","money","money","pct"]
        int_kinds = ["text","int","int","int","int","int","pct"]
        for key, label, pct_key, kinds in [
            ("gmv", "GMV", "WoW", money_kinds),
            ("bookings", "Bookings", "WoW", int_kinds),
            ("receivable", "Receivable", "Ratio", money_kinds),
            ("net_revenue", "Net Revenue", "Ratio", money_kinds),
        ]:
            d = s[key]
            data.append(fmt_row(
                [label, d.get("Flight"), d.get("Tour"), d.get("Visa"),
                 d.get("Hotel"), d.get("Total"), d.get(pct_key)],
                kinds,
            ))
        table = Table(data, repeatRows=1)

    else:
        d = payload["data"]
        data = [["Product","CSS GMV","Retail GMV","Total GMV","Receivable","Bookings"]]
        branch_kinds = ["text","money","money","money","money","int"]
        for prod in ["Flight", "Tour", "Visa", "Hotel", "Vouchers", "Others"]:
            x = d["products"][prod]
            data.append(fmt_row(
                [prod, x["css"], x["retail"], x["gmv"], x["receivable"], x["bookings"]],
                branch_kinds,
            ))
        table = Table(data, repeatRows=1)

    table.setStyle(TableStyle([
        ("BACKGROUND",(0,0),(-1,0),colors.HexColor("#1D4DA2")),
        ("TEXTCOLOR",(0,0),(-1,0),colors.white),
        ("FONTNAME",(0,0),(-1,0),"Helvetica-Bold"),
        ("GRID",(0,0),(-1,-1),0.4,colors.HexColor("#DDE3EA")),
        ("FONTSIZE",(0,0),(-1,-1),8),
        ("VALIGN",(0,0),(-1,-1),"MIDDLE"),
        ("ROWBACKGROUNDS",(0,1),(-1,-1),[colors.white,colors.HexColor("#F6F8FB")]),
        ("LEFTPADDING",(0,0),(-1,-1),5),
        ("RIGHTPADDING",(0,0),(-1,-1),5),
    ]))
    story.append(table)
    story.append(Spacer(1,10))
    story.append(Paragraph("Currency: USD", tiny))
    story.append(Paragraph(f"Source: {SOURCE.source_name} • Last refresh: {SOURCE.last_refresh.isoformat() if SOURCE.last_refresh else 'unknown'}", tiny))
    doc.build(story)
    return bio.getvalue()


def snapshot_base_name(snapshot_date: date) -> str:
    return snapshot_date.strftime("%Y-%m")


def create_month_end_snapshot(snapshot_date: Optional[date] = None) -> dict:
    d = snapshot_date or (now_bdt().date())
    if d != month_end(d):
        raise ValueError("Snapshot date must be the last day of its month.")

    key = snapshot_base_name(d)
    out_dir = SNAPSHOT_DIR / key
    out_dir.mkdir(parents=True, exist_ok=True)

    # Dashboard locked state + full month branch/WBR reports.
    dashboard_payload = dashboard_h2()

    snapshot_meta = {
        "snapshot_date": d.isoformat(),
        "captured_at": now_bdt().isoformat(),
        "source": SOURCE.source_name,
        "source_last_refresh": SOURCE.last_refresh.isoformat() if SOURCE.last_refresh else None,
        "currency": "USD",
        "pages": ["dashboard","banani","chattogram","hq","motijheel","wbr"],
    }
    (out_dir / "snapshot.json").write_text(json.dumps(snapshot_meta, indent=2), encoding="utf-8")
    (out_dir / "dashboard.json").write_text(json.dumps(dashboard_payload, default=str, indent=2), encoding="utf-8")

    # Produce user-facing files for the close.
    (out_dir / "Dashboard.xlsx").write_bytes(export_xlsx("dashboard", None, None))
    (out_dir / "Dashboard.pdf").write_bytes(export_pdf("dashboard", None, None))

    for page in ["banani","chattogram","hq","motijheel"]:
        (out_dir / f"{page}.xlsx").write_bytes(
            export_xlsx(page, d.replace(day=1).isoformat(), d.isoformat())
        )
        (out_dir / f"{page}.pdf").write_bytes(
            export_pdf(page, d.replace(day=1).isoformat(), d.isoformat())
        )

    (out_dir / "wbr.xlsx").write_bytes(
        export_xlsx("wbr", d.isoformat(), d.isoformat())
    )
    (out_dir / "wbr.pdf").write_bytes(
        export_pdf("wbr", d.isoformat(), d.isoformat())
    )

    return snapshot_meta | {"path": str(out_dir)}


def list_snapshots():
    items=[]
    for d in sorted(SNAPSHOT_DIR.iterdir() if SNAPSHOT_DIR.exists() else []):
        if d.is_dir() and (d/"snapshot.json").exists():
            try:
                meta=json.loads((d/"snapshot.json").read_text(encoding="utf-8"))
                items.append(meta)
            except Exception:
                pass
    return items


def maybe_month_end_snapshot():
    now = now_bdt()
    last = month_end(now.date())
    if now.date() != last:
        return
    if not (now.hour == SNAPSHOT_HOUR and now.minute >= SNAPSHOT_MINUTE and now.minute <= SNAPSHOT_MINUTE + 1):
        return
    path = SNAPSHOT_DIR / snapshot_base_name(last) / "snapshot.json"
    if not path.exists():
        try:
            create_month_end_snapshot(last)
        except Exception as exc:
            SOURCE.last_error = f"Month-end snapshot failed: {exc}"


# ============================================================
# API
# ============================================================

@asynccontextmanager
async def lifespan(app: FastAPI):
    refresh_task = asyncio.create_task(auto_refresh_loop())
    snapshot_task = asyncio.create_task(month_end_snapshot_loop())
    yield
    refresh_task.cancel()
    snapshot_task.cancel()
    for task in (refresh_task, snapshot_task):
        try:
            await task
        except asyncio.CancelledError:
            pass


async def auto_refresh_loop():
    # Sleep FIRST: SourceManager.__init__ already refreshed at import time.
    # Refreshing again here delayed the port bind by ~2 minutes on boot.
    while True:
        await asyncio.sleep(REFRESH_MINUTES * 60)
        try:
            # SOURCE.refresh() is blocking requests I/O. Running it directly
            # in the event loop stalls uvicorn, including startup and every
            # in-flight request.
            await asyncio.to_thread(SOURCE.refresh)
        except Exception as exc:
            SOURCE.last_error = str(exc)


async def month_end_snapshot_loop():
    # Check once per minute so a live instance can capture the month-end
    # state at approximately 23:55 Asia/Dhaka rather than depending on
    # the 30-minute data-refresh cadence.
    while True:
        try:
            await asyncio.to_thread(maybe_month_end_snapshot)
        except Exception as exc:
            SOURCE.last_error = f"Month-end snapshot check failed: {exc}"
        await asyncio.sleep(60)


app = FastAPI(
    title="GoZayaan Retail Dashboard",
    lifespan=lifespan,
)

app.mount(
    "/static",
    StaticFiles(directory=str(BASE_DIR / "static")),
    name="static",
)


@app.get("/")
def index():
    return FileResponse(BASE_DIR / "static" / "index.html")


@app.get("/api/health")
def health():
    live_ok = SOURCE.source is not None
    return {
        "ok": live_ok,
        "degraded": SOURCE.degraded,
        "source": SOURCE.source_name,
        "live_required": LIVE_REQUIRED,
        "last_refresh": SOURCE.last_refresh.isoformat()
        if SOURCE.last_refresh else None,
        "last_error": SOURCE.last_error,
        "refresh_minutes": REFRESH_MINUTES,
        "currency": "USD",
        "source_api": bool(APPS_SCRIPT_URL),
    }


@app.post("/api/refresh")
def refresh():
    started = time.perf_counter()

    try:
        SOURCE.refresh()
    except Exception as exc:
        raise HTTPException(
            status_code=502,
            detail=str(exc),
        )

    return {
        "ok": True,
        "source": SOURCE.source_name,
        "refreshed_at": SOURCE.last_refresh.isoformat()
        if SOURCE.last_refresh else None,
        "seconds": round(
            time.perf_counter() - started,
            2
        ),
    }


@app.get("/api/source")
def api_source():
    return {
        "source": SOURCE.source_name,
        "degraded": SOURCE.degraded,
        "live_required": LIVE_REQUIRED,
        "apps_script_url_configured": bool(APPS_SCRIPT_URL),
        "last_refresh": SOURCE.last_refresh.isoformat() if SOURCE.last_refresh else None,
        "last_error": SOURCE.last_error,
    }


@app.get("/api/dashboard")
def api_dashboard():
    data = dashboard_h2()
    return JSONResponse({
        **data,
        "source": SOURCE.source_name,
        "last_refresh": SOURCE.last_refresh.isoformat()
        if SOURCE.last_refresh else None,
    })


@app.get("/api/h2")
def api_h2():
    return {
        "page": "h2",
        "currency": "USD",
        "rows": h2_performance(),
        "source": SOURCE.source_name,
        "last_refresh": SOURCE.last_refresh.isoformat()
        if SOURCE.last_refresh else None,
    }


@app.get("/api/daily-gmv")
def api_daily_gmv(start: str, end: str):
    try:
        start_d = parse_date(start)
        end_d = parse_date(end)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc))

    if start_d > end_d:
        raise HTTPException(
            status_code=400,
            detail="Start date cannot be after end date.",
        )

    if (end_d - start_d).days > 366:
        raise HTTPException(
            status_code=400,
            detail="Range too large — limit to 366 days.",
        )

    rows = daily_company_gmv(start_d, end_d)
    total = sum(r["total"] for r in rows)
    days = len(rows)

    return {
        "page": "daily-gmv",
        "currency": "USD",
        "start": start_d.isoformat(),
        "end": end_d.isoformat(),
        "days": days,
        "rows": rows,
        "total_gmv": total,
        "average_daily_gmv": (total / days) if days else 0.0,
        "source": SOURCE.source_name,
        "last_refresh": SOURCE.last_refresh.isoformat()
        if SOURCE.last_refresh else None,
    }


@app.get("/api/branch/{branch}")
def api_branch(
    branch: str,
    start: str,
    end: str,
):
    canonical = {
        "banani": "Banani",
        "chattogram": "Chattogram",
        "hq": "HQ",
        "motijheel": "Motijheel",
    }.get(branch.lower())

    if not canonical:
        raise HTTPException(
            status_code=404,
            detail="Unknown branch.",
        )

    try:
        start_d = parse_date(start)
        end_d = parse_date(end)
    except ValueError as exc:
        raise HTTPException(
            status_code=400,
            detail=str(exc),
        )

    if start_d > end_d:
        raise HTTPException(
            status_code=400,
            detail="Start date cannot be after end date.",
        )

    data = branch_period(
        canonical,
        start_d,
        end_d,
    )

    # Operational/month progress intentionally uses the selected end date.
    progress = branch_progress(
        canonical,
        end_d,
    )

    # Product achievement bars are locked to the current Bangladesh calendar month
    # and therefore do not move when the report date filter changes.
    today_bdt = datetime.now(BDT).date()
    locked_month_start = month_start(today_bdt)
    locked_month_progress = branch_period(
        canonical,
        locked_month_start,
        today_bdt,
    )

    return {
        "page": "branch",
        "branch": canonical,
        "currency": "USD",
        "start": start,
        "end": end,
        "total_gmv": data["total_gmv"],
        "products": data["products"],
        "progress": {
            "mtd_gmv": progress["total_gmv"],
            "projected": progress["projected"],
            "current_run_rate": progress["current_run_rate"],
            "required_run_rate": progress["required_run_rate"],
            "gap": progress["gap"],
            "active_css": progress["active_css"],
            "pipeline_number": progress["pipeline_number"],
            "pipeline_worth": progress["pipeline_worth"],
            "footfall": data["footfall"],
            "footfall_avg_daily": data["footfall_avg_daily"],
            "targets": progress["targets"],
            "targets_available": progress.get("targets_available", True),
        },
        "locked_month_progress": {
            "month": today_bdt.strftime("%Y-%m"),
            "products": locked_month_progress["products"],
            "targets": locked_month_progress["targets"],
        },
        "source": SOURCE.source_name,
        "last_refresh": SOURCE.last_refresh.isoformat()
        if SOURCE.last_refresh else None,
    }



@app.get("/api/export/{fmt}")
def api_export(
    fmt: str,
    page: str = "dashboard",
    start: Optional[str] = None,
    end: Optional[str] = None,
):
    if fmt not in {"xlsx", "pdf"}:
        raise HTTPException(status_code=400, detail="Format must be xlsx or pdf.")
    try:
        payload = export_xlsx(page, start, end) if fmt == "xlsx" else export_pdf(page, start, end)
    except Exception as exc:
        raise HTTPException(status_code=400, detail=str(exc))

    filename = safe_filename(
        f"{report_title(page, start, end)}.{fmt}"
    )
    media = (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        if fmt == "xlsx"
        else "application/pdf"
    )
    return StreamingResponse(
        io.BytesIO(payload),
        media_type=media,
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@app.get("/api/snapshots")
def api_snapshots():
    return {"snapshots": list_snapshots()}


@app.post("/api/snapshots/month-end")
def api_create_month_end_snapshot(snapshot_date: Optional[str] = None):
    try:
        d = parse_date(snapshot_date) if snapshot_date else now_bdt().date()
        result = create_month_end_snapshot(d)
        return result
    except Exception as exc:
        raise HTTPException(status_code=400, detail=str(exc))


@app.get("/api/snapshots/{month}/{filename}")
def api_snapshot_file(month: str, filename: str):
    if not re.fullmatch(r"\d{4}-\d{2}", month):
        raise HTTPException(status_code=400, detail="Invalid month.")
    safe = safe_filename(filename)
    path = SNAPSHOT_DIR / month / safe
    if not path.exists() or not path.is_file():
        raise HTTPException(status_code=404, detail="Snapshot file not found.")
    return FileResponse(path, filename=safe)


@app.get("/api/wbr")
def api_wbr(week_end: Optional[str] = None):
    """
    week_end omitted -> the currently locked review week (previous
    completed Mon-Sun week; see locked_wbr_sunday). Passing week_end
    still works for reviewing a past week manually — any date is
    snapped to the Sunday of its own week, so "locked to a Sunday"
    holds either way.
    """
    manual_override = week_end is not None

    if manual_override:
        try:
            end = parse_date(week_end)
        except ValueError as exc:
            raise HTTPException(
                status_code=400,
                detail=str(exc),
            )
    else:
        end = locked_wbr_sunday()

    locked_end = locked_wbr_sunday()
    rollover_at = next_wbr_rollover()

    return {
        **wbr_metrics(end),
        "page": "wbr",
        "currency": "USD",
        "manual_override": manual_override,
        "is_current_locked_week": end == locked_end,
        "locked_week_ending": locked_end.isoformat(),
        "next_rollover_at": rollover_at.isoformat(),
        "source": SOURCE.source_name,
        "last_refresh": SOURCE.last_refresh.isoformat()
        if SOURCE.last_refresh else None,
    }


if __name__ == "__main__":
    uvicorn.run(
        "app:app",
        host="0.0.0.0",
        port=int(os.getenv("PORT", "8000")),
        reload=False,
    )
